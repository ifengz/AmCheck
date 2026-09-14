"""review_import —— 测评订单表格(xlsx/xls)解析。

只认"有评价链接的行":没有链接、或链接不是 Amazon 评价 permalink 的行一律丢弃
(需求:以有评价链接为首要加入条件)。同一文件内重复的评价链接也只保留第一条。

列定位:先按表头名匹配(容错各种叫法),匹配不到再退回固定列位
A=刷单编号 / D=产品型号 / F=订单号 / M=评论URL —— 与业务方现有表格一致。
"""

from __future__ import annotations

import io
import re
from pathlib import Path

from engine import parse_link

# 表头别名 → 内部字段。全部小写、去空格后比较。
HEADER_ALIASES = {
    "order_ref": ("刷单编号", "刷单号", "测评编号", "订单编号"),
    "order_no": ("订单号", "amazon订单号", "亚马逊订单号"),
    "model": ("产品型号", "型号", "sku", "产品sku"),
    "url": ("评论url", "测评链接", "评论链接", "评价链接", "评论地址",
            "测评地址", "reviewurl", "review_url", "评价url"),
}

# 兜底列位(0 基):A / D / F / M
FALLBACK_COLS = {"order_ref": 0, "model": 3, "order_no": 5, "url": 12}


class _XlsCell:
    """把 xlrd 的裸值包成 openpyxl 的 cell.value 形态。"""
    __slots__ = ("value",)

    def __init__(self, value):
        self.value = value


class _XlsSheet:
    """xlrd 工作表 → openpyxl 工作表的最小适配器。

    只实现本模块用到的那几个成员(max_row / max_column / cell / title),
    这样下面的解析逻辑对两种引擎完全一致,不用写两遍。
    """

    def __init__(self, sheet):
        self._s = sheet

    @property
    def title(self):
        return self._s.name

    @property
    def max_row(self):
        return self._s.nrows

    @property
    def max_column(self):
        return self._s.ncols

    def cell(self, row, column):
        return _XlsCell(self._s.cell_value(row - 1, column - 1))


class _XlsBook:
    def __init__(self, book):
        self._b = book
        self.worksheets = [_XlsSheet(s) for s in book.sheets()]

    def close(self):
        try:
            self._b.release_resources()
        except Exception:
            pass


# 旧版 .xls(BIFF)是 OLE2 复合文档,魔数与 .xlsx 的 ZIP(PK)完全不同 ——
# 靠魔数分流比靠异常类型判断可靠(openpyxl 对 .xls 抛的是 BadZipFile 而非
# InvalidFileException,按异常名判断会漏)。
OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
ZIP_MAGIC = b"PK"


def _load_workbook(source):
    """打开表格:按魔数分流 xlsx/xlsm(ZIP → openpyxl)与 xls(OLE2 → xlrd)。

    读不出来一律抛 ValueError,消息直接可展示给用户。
    """
    if isinstance(source, (bytes, bytearray)):
        raw = bytes(source)
    elif hasattr(source, "read"):
        raw = source.read()
    else:
        raw = Path(source).read_bytes()

    if raw[:8] == OLE2_MAGIC:
        try:
            import xlrd
        except ImportError:
            raise ValueError(
                "这是旧版 .xls 格式(Excel 97-2003)。请在 Excel 里"
                "「另存为 .xlsx」后重新导入。") from None
        try:
            return _XlsBook(xlrd.open_workbook(file_contents=raw))
        except Exception as e:
            raise ValueError(f"无法读取 .xls 文件:{e}") from e

    if raw[:2] != ZIP_MAGIC:
        raise ValueError(
            "这不是 Excel 表格文件。请选择 .xlsx / .xlsm(.xls 也可,但建议"
            "先另存为 .xlsx)后重试。")

    import openpyxl
    try:
        return openpyxl.load_workbook(io.BytesIO(raw), data_only=True,
                                      read_only=True)
    except Exception as e:
        raise ValueError(f"无法读取表格文件:{e}") from e


def _norm(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", "", str(v)).strip().lower()


def _cell(ws, row: int, col: int) -> str:
    """取单元格文本。数字去掉 Excel 浮点尾巴(20321.0 → 20321)。"""
    v = ws.cell(row=row, column=col + 1).value
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _find_header(ws) -> tuple[int, dict]:
    """在前 12 行里找表头行,返回 (行号 1 基, {字段: 列号 0 基})。找不到返回 (0, {})。"""
    best_row, best_map = 0, {}
    for r in range(1, min(ws.max_row, 12) + 1):
        mapping = {}
        for c in range(ws.max_column):
            head = _norm(ws.cell(row=r, column=c + 1).value)
            if not head:
                continue
            for field, aliases in HEADER_ALIASES.items():
                if field in mapping:
                    continue
                if head in aliases or any(a in head for a in aliases if len(a) >= 3):
                    mapping[field] = c
                    break
        # 至少认出链接列 + 一个业务列才算表头
        if "url" in mapping and len(mapping) >= 2 and len(mapping) > len(best_map):
            best_row, best_map = r, mapping
    return best_row, best_map


def parse_order_file(source) -> tuple[list[dict], dict]:
    """解析表格。

    source: 文件路径 / bytes / 文件对象(如 Starlette UploadFile 读出的 bytes)皆可。
    返回 (rows, stats):
      rows  [{review_id, domain, url, order_ref, order_no, model}]
      stats {total, valid, no_link, dup, header_row, sheet}

    读不出文件(旧版 .xls 且没装 xlrd、文件损坏等)抛 ValueError,消息可直接展示。
    """
    wb = _load_workbook(source)

    ws = wb.worksheets[0]
    header_row, mapping = _find_header(ws)
    if not mapping:
        mapping = dict(FALLBACK_COLS)
        first_data = 1                     # 没表头 → 第 1 行就是数据
    else:
        first_data = header_row + 1

    rows, seen = [], set()
    stats = {"total": 0, "valid": 0, "no_link": 0, "dup": 0,
             "header_row": header_row, "sheet": ws.title}

    def _get(r, field):
        c = mapping.get(field)
        return _cell(ws, r, c) if c is not None else ""

    for r in range(first_data, ws.max_row + 1):
        url_raw = _get(r, "url")
        order_ref = _get(r, "order_ref")
        order_no = _get(r, "order_no")
        model = _get(r, "model")
        if not any((url_raw, order_ref, order_no, model)):
            continue                        # 整行空,跳过
        stats["total"] += 1
        ref = parse_link(url_raw) if url_raw else None
        if ref is None:
            stats["no_link"] += 1           # 无链接 / 链接不合法 → 不加进来
            continue
        key = (ref.domain, ref.review_id)
        if key in seen:
            stats["dup"] += 1
            continue
        seen.add(key)
        stats["valid"] += 1
        rows.append({
            "review_id": ref.review_id,
            "domain": ref.domain,
            "url": ref.url,
            "order_ref": order_ref,
            "order_no": order_no,
            "model": model,
        })

    wb.close()
    return rows, stats


def summarize(stats: dict) -> str:
    """导入弹窗/通知用的一句话统计。"""
    return (f"共 {stats['total']} 行 · 有效 {stats['valid']} 条 · "
            f"无链接跳过 {stats['no_link']} 条 · 文件内重复 {stats['dup']} 条")
